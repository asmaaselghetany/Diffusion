#!/usr/bin/env python3
"""Build thesis progress PPTX from presentation snapshot + decode viz frames."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs' / 'extension_guides' / 'DLMA_Asmaa [Autosaved].pptx'
ALSO = ROOT / 'docs' / 'extension_guides' / 'DLMA_BlockDiffusion_Progress.pptx'
FRAME_DIR = (
    ROOT / 'docs' / 'research' / 'decode_viz'
    / 'compare_baseline_masked_vs_baseline_uniform')

# Visual system: deep slate + teal (avoid purple / cream–terracotta defaults)
NAVY = RGBColor(0x0F, 0x2C, 0x3C)
TEAL = RGBColor(0x1A, 0x6B, 0x5A)
TEAL_SOFT = RGBColor(0xD6, 0xEB, 0xE4)
INK = RGBColor(0x1C, 0x24, 0x28)
MUTED = RGBColor(0x5A, 0x6A, 0x72)
LINE = RGBColor(0xC5, 0xD0, 0xD4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WARN = RGBColor(0x8B, 0x45, 0x13)
OK = RGBColor(0x1A, 0x6B, 0x5A)


def _set_run(run, *, size=18, bold=False, color=INK, name='Calibri'):
  run.font.size = Pt(size)
  run.font.bold = bold
  run.font.color.rgb = color
  run.font.name = name


def _add_textbox(slide, left, top, width, height, text, *,
                 size=18, bold=False, color=INK, align=PP_ALIGN.LEFT,
                 font='Calibri'):
  box = slide.shapes.add_textbox(left, top, width, height)
  tf = box.text_frame
  tf.word_wrap = True
  p = tf.paragraphs[0]
  p.alignment = align
  run = p.add_run()
  run.text = text
  _set_run(run, size=size, bold=bold, color=color, name=font)
  return box


def _add_bullets(slide, left, top, width, height, items, *, size=16,
                 color=INK, level_extra=None):
  box = slide.shapes.add_textbox(left, top, width, height)
  tf = box.text_frame
  tf.word_wrap = True
  for i, item in enumerate(items):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.level = 0
    if level_extra and i in level_extra:
      p.level = level_extra[i]
    p.space_after = Pt(6)
    run = p.add_run()
    run.text = item
    _set_run(run, size=size, color=color)
  return box


def _bar(slide, left, top, width, height, fill=TEAL):
  shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
  shape.fill.solid()
  shape.fill.fore_color.rgb = fill
  shape.line.fill.background()
  return shape


def _footer(slide, page: int, total: int):
  _add_textbox(
      slide, Inches(0.5), Inches(7.05), Inches(8), Inches(0.3),
      'Discrete block diffusion · thesis progress · 2026-09-18',
      size=10, color=MUTED)
  _add_textbox(
      slide, Inches(11.5), Inches(7.05), Inches(1.3), Inches(0.3),
      f'{page} / {total}', size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def _title_slide(prs, title, subtitle, eyebrow='TUM · Discrete Diffusion'):
  slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
  _bar(slide, 0, 0, prs.slide_width, Inches(0.12), TEAL)
  _bar(slide, 0, Inches(6.9), prs.slide_width, Inches(0.5), NAVY)
  _add_textbox(slide, Inches(0.7), Inches(1.8), Inches(11.5), Inches(0.4),
               eyebrow, size=14, bold=True, color=TEAL)
  _add_textbox(slide, Inches(0.7), Inches(2.3), Inches(11.5), Inches(1.5),
               title, size=36, bold=True, color=NAVY)
  _add_textbox(slide, Inches(0.7), Inches(4.2), Inches(11.5), Inches(1.2),
               subtitle, size=18, color=MUTED)
  _add_textbox(slide, Inches(0.7), Inches(7.0), Inches(11), Inches(0.3),
               'Asmaa Elsayed · scoreboard audited 2026-09-18',
               size=12, color=WHITE)
  return slide


def _section(prs, number, title, blurb, page, total):
  slide = prs.slides.add_slide(prs.slide_layouts[6])
  _bar(slide, 0, 0, Inches(0.18), prs.slide_height, TEAL)
  _add_textbox(slide, Inches(0.7), Inches(2.4), Inches(11), Inches(0.4),
               f'SECTION {number}', size=14, bold=True, color=TEAL)
  _add_textbox(slide, Inches(0.7), Inches(2.9), Inches(11), Inches(1.2),
               title, size=32, bold=True, color=NAVY)
  _add_textbox(slide, Inches(0.7), Inches(4.3), Inches(11), Inches(1.2),
               blurb, size=18, color=MUTED)
  _footer(slide, page, total)
  return slide


def _content(prs, title, page, total):
  slide = prs.slides.add_slide(prs.slide_layouts[6])
  _bar(slide, 0, 0, prs.slide_width, Inches(0.08), TEAL)
  _add_textbox(slide, Inches(0.55), Inches(0.25), Inches(12), Inches(0.55),
               title, size=24, bold=True, color=NAVY)
  _bar(slide, Inches(0.55), Inches(0.85), Inches(1.2), Inches(0.05), TEAL)
  _footer(slide, page, total)
  return slide


def _table(slide, left, top, width, rows, col_widths=None, header=True):
  """rows: list[list[str]]."""
  n_rows = len(rows)
  n_cols = len(rows[0])
  height = Inches(0.32 * n_rows + 0.05)
  table = slide.shapes.add_table(
      n_rows, n_cols, left, top, width, height).table
  if col_widths:
    for i, w in enumerate(col_widths):
      table.columns[i].width = w
  for r, row in enumerate(rows):
    for c, val in enumerate(row):
      cell = table.cell(r, c)
      cell.text = str(val)
      for p in cell.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT
        for run in p.runs:
          _set_run(
              run, size=11 if n_cols > 5 else 12,
              bold=(header and r == 0),
              color=WHITE if (header and r == 0) else INK)
      # fill
      fill = cell.fill
      fill.solid()
      if header and r == 0:
        fill.fore_color.rgb = NAVY
      elif r % 2 == 1:
        fill.fore_color.rgb = TEAL_SOFT
      else:
        fill.fore_color.rgb = WHITE
  return table


def _card(slide, left, top, width, height, title, body, accent=TEAL):
  shape = slide.shapes.add_shape(
      MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
  shape.fill.solid()
  shape.fill.fore_color.rgb = WHITE
  shape.line.color.rgb = LINE
  shape.line.width = Pt(1)
  _bar(slide, left, top, Inches(0.08), height, accent)
  _add_textbox(slide, left + Inches(0.2), top + Inches(0.12),
               width - Inches(0.3), Inches(0.35),
               title, size=14, bold=True, color=NAVY)
  _add_textbox(slide, left + Inches(0.2), top + Inches(0.5),
               width - Inches(0.3), height - Inches(0.6),
               body, size=12, color=MUTED)


def build():
  prs = Presentation()
  prs.slide_width = Inches(13.333)
  prs.slide_height = Inches(7.5)

  # Estimate pages for footers (fixed plan)
  TOTAL = 25
  page = 0

  def next_page():
    nonlocal page
    page += 1
    return page

  # 1 Title
  _title_slide(
      prs,
      'From AR Instruct to Block Discrete Diffusion',
      'Paradigms × components × corruption — what actually moves '
      'generation quality on Qwen2.5-1.5B\n'
      'Framed against Fast-dLLM · BlockGen · Unifusion (not full recreations)',
  )
  next_page()

  # 2 Agenda
  s = _content(prs, 'Agenda', next_page(), TOTAL)
  _add_bullets(s, Inches(0.7), Inches(1.2), Inches(11.5), Inches(5.5), [
      '01  Research question & design-space framing',
      '02  Related work map — what we borrow / what we refuse to mix',
      '03  Four axes in depth — paradigm · components · corruption · geometry',
      '04  Why these axes (and not a “full recipe” blob)',
      '05  Results tables — paradigms, ablations, corruption cross-cut',
      '06  Decode geometry — masked vs uniform highlight video',
      '07  Claim K (editability) & hygiene caveats',
      '08  Experiment timeline — done · running · still to run',
  ], size=18)

  # 3 RQ
  s = _content(prs, 'Research questions', next_page(), TOTAL)
  _card(s, Inches(0.55), Inches(1.15), Inches(6.0), Inches(2.4),
        'RQ1 — Conversion',
        'Which AR→block choices (init, Fast-dLLM-style hooks, optional joint AR) '
        'make instruct-scale generation usable after a matched post-train budget?')
  _card(s, Inches(6.75), Inches(1.15), Inches(6.0), Inches(2.4),
        'RQ2 — Fair AR column',
        'After the same data/steps, does block conversion beat matched causal '
        'AR SFT on task quality and parallel decode?')
  _card(s, Inches(0.55), Inches(3.75), Inches(6.0), Inches(2.4),
        'RQ3 — Confounds',
        'Do Fast-dLLM / NLD / BlockGen knobs get confounded with native '
        'geometry, corruption, or decode? Keep families on separate lines.')
  _card(s, Inches(6.75), Inches(3.75), Inches(6.0), Inches(2.4),
        'RQ-B / Claim K',
        'Same AR→block skeleton: does masked vs uniform change generation '
        'quality or only likelihood ranking? Editability as mechanism.')

  # 4 Related work
  s = _content(prs, 'Related work as a design space (not full recreate)', next_page(), TOTAL)
  _table(s, Inches(0.55), Inches(1.2), Inches(12.2), [
      ['Paper / line', 'We take', 'We do not claim'],
      ['Fast-dLLM', 'AR→block + masked components\n(shift, complementary, schedule, CE, hub)',
       'Bit-identical Hub parity as the thesis goal'],
      ['BlockGen', 'Native geometry ideas\n(mixture, ARPC, CE@1)',
       'xfer_* as “BlockGen-uniform” on AR→block'],
      ['Unifusion', 'shift-x₀ probe · GenPPL+H hygiene',
       'GenPPL as fluency / Unifusion peer row'],
      ['Set Diffusion (style)', 'Decode highlight viz vocabulary',
       'Their full training recipe'],
  ], col_widths=[Inches(2.4), Inches(5.0), Inches(4.8)])
  _add_textbox(
      s, Inches(0.55), Inches(5.6), Inches(12), Inches(0.8),
      'Rule: never collapse paradigm + components + corruption into one “full recipe” cell.',
      size=14, bold=True, color=TEAL)

  # 5 Axes overview
  s = _content(prs, 'Four orthogonal axes', next_page(), TOTAL)
  cards = [
      ('Paradigm', 'Training path / objective\nNative · AR→block · AR→full · Joint'),
      ('Components', 'Loss / mask / attn / schedule knobs\nPortable across corruption'),
      ('Corruption', 'Forward process\nMasked · Uniform · Hybrid'),
      ('Geometry', 'Block sizes\nFixed-32 · mixture / weights'),
  ]
  for i, (t, b) in enumerate(cards):
    _card(s, Inches(0.45 + i * 3.2), Inches(1.4), Inches(3.0), Inches(3.6),
          t, b)
  _add_textbox(
      s, Inches(0.55), Inches(5.4), Inches(12), Inches(1.0),
      'Eval is a fifth orthogonal layer: DualCache / hubmatch thr=1 vs baseline ancestral; '
      'likelihood MMLU vs generative tasks. GenPPL + entropy = collapse hygiene only.',
      size=14, color=MUTED)

  # Section divider
  _section(prs, '01', 'Axis depth — Paradigm',
           'What training path are we on? Everything else hangs off this choice.',
           next_page(), TOTAL)

  # 6 Paradigm depth
  s = _content(prs, 'Paradigm — why these four paths', next_page(), TOTAL)
  _table(s, Inches(0.45), Inches(1.15), Inches(12.4), [
      ['Paradigm', 'Definition', 'Why include'],
      ['Native', 'Train block diffusion from scratch',
       'Floor: conversion must beat this; BlockGen home'],
      ['AR → block', 'Convert Instruct LM → fixed-block diffusion',
       'Main thesis path (Fast-dLLM-style conversion)'],
      ['AR → full sequence', 'Matched causal AR SFT (no blocks)',
       'Fair AR column for Tab 1 / RQ2'],
      ['Joint', 'Block diffusion + AR auxiliary loss',
       'NLD-style bridge without abandoning blocks'],
  ], col_widths=[Inches(2.5), Inches(4.7), Inches(5.2)])
  _add_textbox(
      s, Inches(0.55), Inches(5.3), Inches(12), Inches(1.0),
      'Components attach mainly to AR→block. They are not paradigms.',
      size=15, bold=True, color=TEAL)

  # 7 Paradigm results
  s = _content(prs, 'Paradigm results — AR→block & floors (ready cells)', next_page(), TOTAL)
  _table(s, Inches(0.35), Inches(1.1), Inches(12.6), [
      ['Run', 'Corruption', 'GSM', 'IFE', 'MMLU', 'GenPPL (H̄)', 'Status'],
      ['Scratch block', 'masked', '0.5', '8.1', '23.0', '—‡', 'Ready'],
      ['Baseline masked', 'masked', '62.1', '20.5', '39.8†', '128 (4.80)', 'Ready'],
      ['Math mix (5 knobs)', 'masked', '66.8', '22.4', '30.6†', '89 (4.76)', 'Ready'],
      ['Chat mix / Mixed SFT', 'masked', '43–60', '23–24', '38–40', '—', 'Ready'],
      ['Baseline uniform', 'uniform', '—', '10.7*', '—', '61 (6.04)', 'IFE*; GSM queued'],
      ['Joint AR masked', 'masked', '38.5', '25.0', '42.0', '—‡', 'Ready'],
      ['Matched AR SFT', '—', '—', '—', '—', '—', 'Queued'],
  ], col_widths=[Inches(2.6), Inches(1.5), Inches(1.2), Inches(1.2),
                 Inches(1.4), Inches(2.3), Inches(2.4)])
  _add_textbox(
      s, Inches(0.45), Inches(5.55), Inches(12.3), Inches(0.9),
      '† hubll MMLU siblings · ‡ do not cite collapsed free-gen GenPPL≈402 · '
      '* U0 IFE 10.7 used FORCE_GREEDY=1 (invalid for uniform); ancestral rescore pending. '
      'GenPPL = first_chunk hygiene, not fluency.',
      size=11, color=MUTED)

  # Section components
  _section(prs, '02', 'Axis depth — Components',
           'Explicit knobs that change loss, masks, attention, or schedule.',
           next_page(), TOTAL)

  # 8 Components why
  s = _content(prs, 'Components — what each does in training', next_page(), TOTAL)
  _table(s, Inches(0.4), Inches(1.1), Inches(12.5), [
      ['Component', 'Config', 'Effect'],
      ['shift', 'shift_loss_targets', 'At i, score clean token i+1 (AR next-token as x₀)'],
      ['complementary', 'complementary_masks', 'Paired masks m and ~m each step'],
      ['mask schedule', 'mask_schedule', 'Remap t → mask rate (e.g. Fast-dLLM)'],
      ['plain CE', 'loss_weighting=plain_ce', 'Unweighted mask-site CE (Hub denom)'],
      ['hub train parity', 'hub_struct_attn_only / …', 'Structural block-attn; pad visibility'],
      ['joint AR', 'joint_ar_alpha', 'Add clean-stream AR loss'],
      ['causal clean', 'causal_clean_stream', 'Token-causal AR pass on clean stream'],
      ['intra-block anneal', 'intra_block_attn_anneal_steps', 'Causal → bidirectional inside block'],
      ['kernel anneal', 'kernel_anneal_steps', 'Hybrid: ramp p(uniform) over training'],
  ], col_widths=[Inches(2.4), Inches(3.6), Inches(6.5)])

  # 9 Ablations
  s = _content(prs, 'Component ablations — AR→block masked', next_page(), TOTAL)
  _table(s, Inches(0.5), Inches(1.15), Inches(12.3), [
      ['Run', 'Components', 'GSM', 'IFE', 'MMLU', 'GenPPL'],
      ['Baseline masked', 'none', '62.1', '20.5', '39.8†', '128'],
      ['+ shift only', 'shift', '64.0', '21.1', '33.7', '88'],
      ['+ complementary only', 'complementary', '53.3', '18.1', '37.3', '108'],
      ['+ shift + complementary', 'both', '50.8', '17.0', '32.0', '107'],
      ['+ intra-block anneal', 'anneal', '61.1', '21.4', '42.7', '134'],
      ['Math mix', 'all five', '66.8', '22.4', '30.6†', '89'],
  ], col_widths=[Inches(3.2), Inches(2.6), Inches(1.4), Inches(1.4),
                 Inches(1.6), Inches(2.1)])
  _add_bullets(s, Inches(0.6), Inches(5.0), Inches(12), Inches(1.6), [
      'shift ↑GSM / ↓GenPPL · complementary alone hurts · shift+comp worse than either alone here',
      'all five + math data → best GSM · intra-block anneal ↑MMLU · joint AR ↑IFE/MMLU but ↓GSM',
  ], size=14)

  # Section corruption
  _section(prs, '03', 'Axis depth — Corruption',
           'Orthogonal forward process: masked absorbing vs uniform redraw vs hybrid.',
           next_page(), TOTAL)

  # 10 Corruption why
  s = _content(prs, 'Corruption — why masked / uniform / hybrid', next_page(), TOTAL)
  _card(s, Inches(0.5), Inches(1.2), Inches(4.0), Inches(4.2),
        'Masked (absorbing)',
        'Unknown state = MASK. Confidence decode (hubmatch) can commit '
        'high-confidence sites. Main Fast-dLLM conversion arm.\n\n'
        'Baseline IFE 20.5 · GSM 62.1')
  _card(s, Inches(4.7), Inches(1.2), Inches(4.0), Inches(4.2),
        'Uniform',
        'Tokens redrawn from Unif(V). Revisable in place — editable-token story '
        '(Claim K). No confidence-unmask loop.\n\n'
        'Never greedy-argmax the posterior (locks prior noise).')
  _card(s, Inches(8.9), Inches(1.2), Inches(4.0), Inches(4.2),
        'Hybrid',
        'MASK + p(uniform). Kernel anneal ramps p. B4 probe — not the main '
        'conversion claim.\n\n'
        'Queued (+ GenPPL hygiene already on disk).')

  # 11 Decode viz
  s = _content(prs, 'Decode geometry — masked vs uniform (same prompt)', next_page(), TOTAL)
  mid = FRAME_DIR / 'slide_frame_mid.png'
  if mid.exists():
    s.shapes.add_picture(str(mid), Inches(0.4), Inches(1.05), width=Inches(12.5))
  _add_textbox(
      s, Inches(0.5), Inches(5.35), Inches(12.3), Inches(1.2),
      'Left: masked + hubmatch thr=1 (MASK → commit).  Right: uniform ancestral '
      '(revise tokens in the gold window). Gold = active block. '
      'Caveat: decode stacks differ — teaser for corruption geometry, not a matched score claim. '
      'Animated GIF: docs/research/decode_viz/compare_…/preview.gif',
      size=12, color=MUTED)

  # 12 Corruption results + hygiene
  s = _content(prs, 'Corruption cross-cut & GenPPL hygiene', next_page(), TOTAL)
  _table(s, Inches(0.5), Inches(1.15), Inches(12.3), [
      ['Contrast', 'Number', 'Read as'],
      ['Masked vs uniform IFE (skeleton)', '20.5 vs 10.7*', 'Quality gap — but *greedy bug on U0'],
      ['Masked GenPPL (H̄)', '128 (4.80)', 'Hygiene baseline'],
      ['Uniform GenPPL (H̄)', '61 (6.04)', 'Lower PPL + higher H ≠ better English'],
      ['Uniform block-ELBO bs1→32', 'PPL 3.8 → 5.3', 'Harder at training block size'],
  ], col_widths=[Inches(4.5), Inches(2.8), Inches(5.0)])
  _add_bullets(s, Inches(0.6), Inches(4.6), Inches(12), Inches(2.0), [
      'Cite GenPPL only with unigram entropy; never as fluency.',
      'U0 ancestral decode is English-ish but weaker than masked hubmatch; soup GIFs were greedy.',
      'Fair next score: U0 IFE with FORCE_GREEDY=0 (job queued).',
  ], size=14)

  # Section geometry
  _section(prs, '04', 'Axis depth — Geometry & Claim K',
           'Block sizes are orthogonal to corruption. Editability is a mechanism claim.',
           next_page(), TOTAL)

  # 13 Geometry
  s = _content(prs, 'Geometry — fixed-32 vs mixture', next_page(), TOTAL)
  _card(s, Inches(0.55), Inches(1.3), Inches(6.0), Inches(3.8),
        'Fixed block-32',
        'Default conversion skeleton (C0 / U0). Matched for Claim K.\n\n'
        'Keeps paradigm/component contrasts clean.')
  _card(s, Inches(6.8), Inches(1.3), Inches(6.0), Inches(3.8),
        'Block-size mix (e.g. 5%×1 / 95%×32)',
        'BlockGen-inspired geometry on conversion line.\n\n'
        'Masked mix: GSM 63.3 · IFE 20.7 — close to baseline.\n'
        'Uniform mix: GenPPL 55 — still hygiene; lm-eval queued.\n\n'
        'Do not confuse mix with “uniform baseline”.')

  # 14 Claim K
  s = _content(prs, 'Claim K — editable tokens (offline package)', next_page(), TOTAL)
  _add_bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(2.2), [
      'Only C0 (masked baseline) vs U0 (uniform baseline) — same math mix / skeleton.',
      'Taxonomy first: label L1/L2/G1/G2 on IFEval fails before causal language.',
      'Go criterion: local share ≥ 30% on C0; then ask whether U0 helps the local slice.',
  ], size=16)
  _table(s, Inches(0.5), Inches(3.6), Inches(12.3), [
      ['Artifact', 'Role'],
      ['C0 / U0 pilot label sheets', 'N=50 fails each'],
      ['CLAIM_K_paired_label_sheet.csv', 'Paired contrast (82 rows)'],
      ['CLAIM_K_export_status.json', 'Strict/loose + family counts'],
  ], col_widths=[Inches(5.5), Inches(6.8)])

  # 15 Why axes
  s = _content(prs, 'Why this factorization (methodological)', next_page(), TOTAL)
  _add_bullets(s, Inches(0.7), Inches(1.2), Inches(12), Inches(5.5), [
      'Separates training path (paradigm) from portable knobs (components).',
      'Keeps corruption orthogonal so masked≠uniform is a clean Claim K contrast.',
      'Puts BlockGen geometry on the right line (native / xfer) — not a fake “BlockGen convert”.',
      'Forces explicit decode profile in tables (hubmatch thr=1 vs baseline ancestral).',
      'Rejects “full recipe” cells that hide which knob moved the number.',
      'Inspired by how Fast-dLLM / BlockGen / Unifusion each own a slice of the space — we map, not merge.',
  ], size=17)

  # Timeline section
  _section(prs, '05', 'Timeline & queue',
           'What landed · what is running · what still needs a number.',
           next_page(), TOTAL)

  # 16 Timeline
  s = _content(prs, 'Experiment timeline (Sep 2026)', next_page(), TOTAL)
  _table(s, Inches(0.4), Inches(1.15), Inches(12.5), [
      ['When', 'Milestone'],
      ['Earlier', 'C0 masked baseline + math/chat mixes; Tab-2 ablations scored'],
      ['Sep 17', 'U0 uniform train finished; hygiene GenPPL dual/NFE; IFE samples (greedy)'],
      ['Sep 18 AM', 'Scoreboard audit; decode-viz pipeline; Claim K export sheets'],
      ['Sep 18 mid', 'Uniform audit: greedy posterior locks noise → sampler + submit fixes'],
      ['Sep 18 PM', 'Ancestral uniform GIF + compare; ancestral U0 IFE submitted'],
      ['Soft ~17:00', 'GSM / MMLU-gen / AR-SFT / uniform knobs / hybrid fills'],
  ], col_widths=[Inches(2.2), Inches(10.3)])

  # 17 Queue status
  s = _content(prs, 'Results status board', next_page(), TOTAL)
  _table(s, Inches(0.4), Inches(1.15), Inches(12.5), [
      ['State', 'Items'],
      ['Ready to cite', 'Native scratch · C0 · math/chat mixes · Tab-2 ablations · '
       'joint masked · block-size mix masked · Claim K sheets'],
      ['Cite with caveat', 'U0 IFE 10.7 (greedy) · GenPPL hygiene pairs · decode GIF '
       '(profile mismatch noted)'],
      ['Running / pending', 'GenPPL hygiene jobs · U0 GSM · ancestral U0 IFE · '
       'matched AR SFT · U0_shift / anneal / hybrid / continue-FT evals'],
      ['Do not cite', 'Free-gen GenPPL≈402 · ARPC as fluency · xfer as matched U0'],
  ], col_widths=[Inches(2.6), Inches(9.9)])

  # 18 Next actions
  s = _content(prs, 'Immediate next actions', next_page(), TOTAL)
  _add_bullets(s, Inches(0.7), Inches(1.3), Inches(12), Inches(5), [
      'Land ancestral U0 IFE — replace 10.7 in the scoreboard if valid.',
      'Label Claim K sheets (50+50) before any causal editability claim.',
      'Fair decode: C0 baseline ancestral vs U0 ancestral (optional GIF).',
      'Fill matched AR SFT column when 1857323 eval returns.',
      'Uniform knobs (shift / anneal): do not promote until generative scores exist '
      '(early free-gen for U0_shift looks worse, not better).',
  ], size=17)

  # 19 Takeaways
  s = _content(prs, 'Takeaways for the talk', next_page(), TOTAL)
  _add_bullets(s, Inches(0.7), Inches(1.3), Inches(12), Inches(5), [
      'Conversion works: scratch 0.5 GSM → C0 62 → math mix 67 under matched budget.',
      'Components are not free: shift helps; complementary alone hurts; interactions matter.',
      'Corruption matters: uniform is a real axis (Claim K), but eval must respect decode rules.',
      'Hygiene ≠ fluency: quote GenPPL with H̄ only.',
      'Keep the map explicit — that is the contribution relative to Fast-dLLM / BlockGen / Unifusion.',
  ], size=17)

  # Thank you
  slide = prs.slides.add_slide(prs.slide_layouts[6])
  _bar(slide, 0, 0, prs.slide_width, Inches(0.12), TEAL)
  _bar(slide, 0, Inches(6.9), prs.slide_width, Inches(0.5), NAVY)
  _add_textbox(slide, Inches(0.7), Inches(2.6), Inches(12), Inches(1),
               'Thank you', size=40, bold=True, color=NAVY)
  _add_textbox(slide, Inches(0.7), Inches(3.8), Inches(12), Inches(1.2),
               'Questions · discussion\n'
               'Scoreboard: docs/research/PRESENTATION_SNAPSHOT_2026-09-18.md\n'
               'Decode GIF: docs/research/decode_viz/compare_…/preview.gif',
               size=16, color=MUTED)
  next_page()

  OUT.parent.mkdir(parents=True, exist_ok=True)
  prs.save(str(OUT))
  prs.save(str(ALSO))
  print('saved', OUT)
  print('saved', ALSO)
  print('slides', len(prs.slides))


if __name__ == '__main__':
  build()
