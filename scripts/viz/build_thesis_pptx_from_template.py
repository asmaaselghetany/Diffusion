#!/usr/bin/env python3
"""Rewrite thesis content into the original DLMA designed slides (no slide cloning)."""

from __future__ import annotations

import shutil
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
BACKUP = ROOT / 'docs' / 'extension_guides' / 'DLMA_Asmaa_medical_BACKUP.pptx'
OUT = ROOT / 'docs' / 'extension_guides' / 'DLMA_Asmaa [Autosaved].pptx'
ALSO = ROOT / 'docs' / 'extension_guides' / 'DLMA_BlockDiffusion_Progress.pptx'
FRAME = (
    ROOT / 'docs' / 'research' / 'decode_viz'
    / 'compare_baseline_masked_vs_baseline_uniform' / 'slide_frame_mid.png')


def set_text(shape, text: str) -> None:
  if not shape.has_text_frame:
    return
  tf = shape.text_frame
  lines = (text or '').split('\n')
  for p in list(tf.paragraphs)[1:]:
    p._p.getparent().remove(p._p)
  p0 = tf.paragraphs[0]
  if p0.runs:
    p0.runs[0].text = lines[0]
    for r in list(p0.runs)[1:]:
      r.text = ''
  else:
    p0.add_run().text = lines[0]
  for line in lines[1:]:
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = line
    if p0.runs and p0.runs[0].font.size:
      r.font.size = p0.runs[0].font.size


def shapes_text(slide):
  out = []
  for sh in slide.shapes:
    if sh.has_text_frame:
      out.append(sh)
  out.sort(key=lambda s: (s.top, s.left))
  return out


def find_text(slide, *needles):
  for sh in slide.shapes:
    if not sh.has_text_frame:
      continue
    t = sh.text_frame.text
    for n in needles:
      if n.lower() in t.lower():
        return sh
  return None


def replace_exact(slide, mapping: dict[str, str]) -> int:
  n = 0
  for sh in slide.shapes:
    if not sh.has_text_frame:
      continue
    t = sh.text_frame.text.strip()
    if t in mapping:
      set_text(sh, mapping[t])
      n += 1
      continue
    for old, new in mapping.items():
      if t == old or t.startswith(old) and len(t) < len(old) + 8:
        set_text(sh, new)
        n += 1
        break
  return n


def clear_long_bodies(slide, keep_shapes, min_len=25):
  for sh in slide.shapes:
    if sh in keep_shapes or not sh.has_text_frame:
      continue
    t = sh.text_frame.text.strip()
    if 'Slide' in t:
      continue
    if len(t) >= min_len:
      set_text(sh, '')


def add_table(slide, rows, left, top, width, height):
  n_rows, n_cols = len(rows), len(rows[0])
  ts = slide.shapes.add_table(n_rows, n_cols, left, top, width, height)
  table = ts.table
  for r, row in enumerate(rows):
    for c, val in enumerate(row):
      cell = table.cell(r, c)
      cell.text = str(val)
      for p in cell.text_frame.paragraphs:
        for run in p.runs:
          run.font.size = Pt(9 if n_cols > 5 else 10)
          run.font.bold = (r == 0)
  return ts


def keep_only(prs, indices: list[int]) -> None:
  sldIdLst = prs.slides._sldIdLst
  all_ids = list(sldIdLst)
  ordered = [all_ids[i] for i in indices]
  for child in list(sldIdLst):
    sldIdLst.remove(child)
  for el in ordered:
    sldIdLst.append(el)


def build():
  shutil.copy2(BACKUP, OUT)
  prs = Presentation(str(OUT))

  # 0 Title
  s = prs.slides[0]
  set_text(shapes_text(s)[0], 'From AR Instruct to\nBlock Discrete Diffusion')
  set_text(
      shapes_text(s)[1],
      'Paradigms × components × corruption — thesis progress\n'
      'Deep Learning for Medical Applications · 2026-09-18')

  # 1 TOC
  s = prs.slides[1]
  replace_exact(s, {
      'IN THIS ISSUE': 'IN THIS TALK',
      'LLMS': 'BLOCK\nDIFFUSION',
      'LLMS in Medical Diagnosis': 'Research questions & framing',
      'Text-Based Approach': 'Four axes in depth',
      'Approach 1': 'Results & ablations',
      'Approach 2': 'Decode viz · Claim K',
      'Approach 3': 'Timeline & queue',
      'Conclusion': 'Takeaways',
      'Medical Diagnosis': 'AR → Block',
      'Text-based': 'Axes',
  })
  for sh in s.shapes:
    if sh.has_text_frame and 'Asmaa' in sh.text_frame.text:
      set_text(sh, 'Asmaa Elsayed · DLMA')

  # 2 Design space (left art)
  s = prs.slides[2]
  h = find_text(s, 'WHAT IS')
  if h:
    set_text(h, 'THE DESIGN SPACE')
  for sh in shapes_text(s):
    if sh is h or 'Slide' in sh.text_frame.text:
      continue
    if len(sh.text_frame.text) > 40:
      set_text(
          sh,
          'We do not recreate Fast-dLLM, BlockGen, or Unifusion end-to-end.\n\n'
          'We map their ideas onto orthogonal axes — paradigm, components, '
          'corruption, geometry — so each number has a named cause.')
      break

  # 3 Motivation → RQ one-liner
  s = prs.slides[3]
  h = find_text(s, 'Motivation')
  if h:
    set_text(
        h,
        'Research question\n\n'
        'Which conversion choices make instruct-scale\n'
        'block diffusion usable — without confounding\n'
        'path, knobs, corruption, and decode?')

  # 6 Three RQs (approaches template)
  s = prs.slides[6]
  replace_exact(s, {
      'Overview of the Three Approaches': 'Three research questions',
      'Approaches': 'RQs',
      'Personalized EHR-based Disease Prediction':
          'RQ1 — Which AR→block choices make generation usable?',
      'Hybrid Models with Structured Data Approach':
          'RQ2 — Does conversion beat matched causal AR SFT?',
      'LLMs with Few-Shot Learning':
          'RQ3 / Claim K — Separate confounds; masked vs uniform',
  })

  # 7 Section: related work
  s = prs.slides[7]
  for sh in shapes_text(s):
    if 'Slide' in sh.text_frame.text:
      continue
    if len(sh.text_frame.text) > 20:
      set_text(
          sh,
          'Related work as a design space\n'
          'Fast-dLLM · BlockGen · Unifusion · Set-Diffusion (viz)')
      break

  # 8 Related work body
  s = prs.slides[8]
  h = find_text(s, 'Problem Statement')
  if h:
    set_text(h, 'What we borrow — and refuse to mix')
  for sh in shapes_text(s):
    if sh is h or 'Slide' in sh.text_frame.text:
      continue
    if len(sh.text_frame.text) > 30:
      set_text(
          sh,
          'Fast-dLLM → AR→block + masked components '
          '(shift, complementary, schedule, plain CE, hub parity).\n\n'
          'BlockGen → native geometry (mixture, ARPC) — never claim xfer = BlockGen.\n\n'
          'Unifusion → shift-x₀ probe + GenPPL+H hygiene — not fluency.\n\n'
          'Set Diffusion → decode highlight style only.\n\n'
          'Rule: no “full recipe” cells that hide which knob moved the score.')
      break

  # 10 Paradigm vs Components (two-box template)
  s = prs.slides[10]
  replace_exact(s, {
      'Retrieval-Augmented Generation (RAG)': 'Two primary axes',
      'Retrieval Component': 'Paradigm',
      'Generation Component': 'Components',
  })
  for sh in s.shapes:
    if not sh.has_text_frame:
      continue
    t = sh.text_frame.text.strip()
    if t.startswith('To fetch'):
      set_text(
          sh,
          'Training path / objective:\nNative · AR→block · AR→full · Joint')
    elif t.startswith('To generate'):
      set_text(
          sh,
          'Portable knobs on AR→block:\n'
          'shift · complementary · schedule · CE · anneals · joint AR')

  # 12 Corruption / Geometry / Eval
  s = prs.slides[12]
  replace_exact(s, {
      'Methodology': 'Orthogonal layers',
      'Data Flow': 'Corruption',
      'Retrieval': 'Geometry',
      'Contextualization': 'Eval',
  })

  # 14 Results header + paradigm table overlay
  s = prs.slides[14]
  h = find_text(s, 'Results')
  if h:
    set_text(h, 'Paradigm results')
  # clear decorative body if any large empty-ish areas — add table
  add_table(
      s,
      [
          ['Run', 'Corr.', 'GSM', 'IFE', 'MMLU', 'GenPPL', 'Status'],
          ['Scratch', 'mask', '0.5', '8.1', '23.0', '—‡', 'Ready'],
          ['Baseline masked', 'mask', '62.1', '20.5', '39.8†', '128', 'Ready'],
          ['Math mix (5 knobs)', 'mask', '66.8', '22.4', '30.6†', '89', 'Ready'],
          ['Baseline uniform', 'unif', '—', '10.7*', '—', '61', 'IFE*; GSM Q'],
          ['Joint AR masked', 'mask', '38.5', '25.0', '42.0', '—‡', 'Ready'],
          ['Matched AR SFT', '—', '—', '—', '—', '—', 'Queued'],
      ],
      Inches(0.35), Inches(1.15), Inches(9.3), Inches(3.6),
  )
  note = s.shapes.add_textbox(Inches(0.35), Inches(4.85), Inches(9.3), Inches(0.5))
  note.text_frame.word_wrap = True
  r = note.text_frame.paragraphs[0].add_run()
  r.text = (
      '† hubll MMLU · ‡ do not cite GenPPL≈402 · * U0 IFE used greedy '
      '(invalid for uniform); ancestral rescore pending · GenPPL = hygiene')
  r.font.size = Pt(9)

  # 15 Highlights (strengths cards)
  s = prs.slides[15]
  h = find_text(s, 'Strengths')
  if h:
    set_text(h, 'Component effects (ready)')
  replace_exact(s, {
      'Personalization of Predictions': 'Scratch→C0 GSM 0.5→62.1',
      'Complex Query Handling': 'Math mix GSM 66.8',
      'Scalability': 'Shift only GSM 64.0 · GenPPL 88',
      'Generalization to Rare Diseases': 'Intra-block anneal MMLU 42.7',
      'Interpretability': 'Joint AR IFE 25.0 · MMLU 42.0',
      'Contextual Accuracy': 'GenPPL = hygiene (cite with H̄)',
  })

  # 16–18 Caveats (criticism template)
  caveats = [
      (16, 'Caveat — uniform IFE',
       'U0 IFE 10.7 used FORCE_GREEDY=1. On uniform, argmax(q) locks prior noise '
       '→ soup. Ancestral decode is English-ish; ancestral IFE pending.',
       'Fix: sampler auto-disables greedy on uniform; FORCE_GREEDY=0 on submit.'),
      (17, 'Caveat — decode compare',
       'Masked GIF = hubmatch thr=1; uniform GIF = baseline ancestral. '
       'Good corruption teaser — not a matched score claim.',
       'Optional fair GIF: both baseline ancestral.'),
      (18, 'Caveat — GenPPL hygiene',
       'Uniform GenPPL≈61 with higher H̄ than masked 128 — collapse hygiene, '
       'not better English. Never cite free-gen GenPPL≈402.',
       'Always quote (PPL, H̄); first_chunk only.'),
  ]
  for idx, title, body, fix in caveats:
    s = prs.slides[idx]
    h = find_text(s, 'Criticism')
    if h:
      set_text(h, title)
    bodies = [sh for sh in shapes_text(s)
              if sh is not h and 'Slide' not in sh.text_frame.text
              and len(sh.text_frame.text.strip()) > 15]
    if len(bodies) >= 1:
      set_text(bodies[0], body)
    if len(bodies) >= 2:
      set_text(bodies[1], fix)
    # clear any remaining long medical leftovers
    for sh in bodies[2:]:
      set_text(sh, '')
    for sh in shapes_text(s):
      t = sh.text_frame.text.lower()
      if any(k in t for k in (
          'ehr', 'rag', 'hospital', 'disease', 'clinical',
          'retrieval module', 'latency', 'patient')):
        if 'fast-dllm' in t or 'blockgen' in t:
          continue
        set_text(sh, '')

  # 19 Why factorization
  s = prs.slides[19]
  h = find_text(s, 'Applicability')
  if h:
    set_text(h, 'Why this factorization')
  for sh in shapes_text(s):
    if sh is h or 'Slide' in sh.text_frame.text:
      continue
    if len(sh.text_frame.text) > 20:
      set_text(
          sh,
          'Separates path (paradigm) from knobs (components).\n'
          'Keeps corruption orthogonal for Claim K.\n'
          'Puts BlockGen geometry on the right line.\n'
          'Forces explicit decode profiles in tables.\n'
          'Rejects “full recipe” cells.')
    # clear other medical leftovers on this slide
  for sh in shapes_text(s):
    t = sh.text_frame.text.lower()
    if any(k in t for k in ('ehr', 'hospital', 'healthcare', 'clinician')):
      if 'separates path' in t:
        continue
      set_text(sh, '')
  # re-apply body if cleared
  h = find_text(s, 'Why this factorization')
  has_body = any(
      'Separates path' in sh.text_frame.text for sh in shapes_text(s))
  if h and not has_body:
    for sh in shapes_text(s):
      if sh is h or 'Slide' in sh.text_frame.text:
        continue
      if sh.top > h.top:
        set_text(
            sh,
            'Separates path (paradigm) from knobs (components).\n'
            'Keeps corruption orthogonal for Claim K.\n'
            'Puts BlockGen geometry on the right line.\n'
            'Forces explicit decode profiles in tables.\n'
            'Rejects “full recipe” cells.')
        break

  # 20 Ablation section title (was Health-LLM paper title slide)
  s = prs.slides[20]
  for sh in shapes_text(s):
    if 'Slide' in sh.text_frame.text:
      continue
    if len(sh.text_frame.text) > 10:
      set_text(sh, 'Component ablations\nAR→block masked')
      break

  # 21 Ablation problem/body → table
  s = prs.slides[21]
  h = find_text(s, 'Problem')
  if h:
    set_text(h, 'What each knob does')
  for sh in shapes_text(s):
    if sh is h or 'Slide' in sh.text_frame.text:
      continue
    if len(sh.text_frame.text) > 30:
      set_text(sh, '')  # make room
  add_table(
      s,
      [
          ['Run', 'Components', 'GSM', 'IFE', 'MMLU', 'GenPPL'],
          ['Baseline', 'none', '62.1', '20.5', '39.8', '128'],
          ['+ shift', 'shift', '64.0', '21.1', '33.7', '88'],
          ['+ complementary', 'comp', '53.3', '18.1', '37.3', '108'],
          ['+ shift+comp', 'both', '50.8', '17.0', '32.0', '107'],
          ['+ intra anneal', 'anneal', '61.1', '21.4', '42.7', '134'],
          ['Math mix', 'all five', '66.8', '22.4', '30.6', '89'],
      ],
      Inches(0.4), Inches(1.35), Inches(9.2), Inches(3.5),
  )

  # 29 Results-compared slide if exists — use for status board (index 29)
  s = prs.slides[29]
  for sh in shapes_text(s):
    if 'Slide' in sh.text_frame.text:
      continue
    if len(sh.text_frame.text) > 5:
      set_text(sh, 'Results status board')
      break
  add_table(
      s,
      [
          ['State', 'Items'],
          ['Ready', 'C0 · math/chat mixes · Tab-2 · joint masked · Claim K sheets'],
          ['Caveat', 'U0 IFE 10.7 (greedy) · GenPPL hygiene · decode GIF profiles'],
          ['Pending', 'Ancestral U0 IFE · U0 GSM · AR SFT · U0 knobs · hybrid'],
          ['Do not cite', 'GenPPL≈402 · ARPC as fluency · xfer as matched U0'],
      ],
      Inches(0.4), Inches(1.3), Inches(9.2), Inches(3.2),
  )

  # 36 Few-shot paper title → Corruption axis
  s = prs.slides[36]
  for sh in shapes_text(s):
    if 'Slide' in sh.text_frame.text:
      continue
    if len(sh.text_frame.text) > 10:
      set_text(sh, 'Corruption axis\nMasked · Uniform · Hybrid')
      break

  # 37 Problem → corruption why
  s = prs.slides[37]
  h = find_text(s, 'Problem')
  if h:
    set_text(h, 'Why these three')
  for sh in shapes_text(s):
    if sh is h or 'Slide' in sh.text_frame.text:
      continue
    if len(sh.text_frame.text) > 30:
      set_text(
          sh,
          'Masked: absorbing MASK + confidence decode (main Fast-dLLM arm).\n\n'
          'Uniform: Unif(V) redraw — revisable tokens (Claim K). Never greedy-argmax.\n\n'
          'Hybrid: MASK + p(uniform); kernel anneal on B4 probe.')
      break

  # 11 already used — for decode viz use slide 46 (workflow) or 47
  # Use slide 14 is results — better: slide 5 (text-based) if mostly empty art
  # Put decode viz on slide 46
  s = prs.slides[46]
  for sh in shapes_text(s):
    t = sh.text_frame.text.strip()
    if 'Workflow' in t or len(t) > 8:
      if 'Slide' in t:
        continue
      set_text(sh, 'Decode geometry — masked vs uniform')
      break
  # clear other text to reduce clutter
  for sh in shapes_text(s):
    t = sh.text_frame.text.strip()
    if t.startswith('Decode geometry'):
      continue
    if 'Slide' in t:
      continue
    if len(t) > 10:
      set_text(sh, '')
  if FRAME.exists():
    prs.slides[46].shapes.add_picture(
        str(FRAME), Inches(0.3), Inches(1.2), width=Inches(9.4))
  cap = prs.slides[46].shapes.add_textbox(
      Inches(0.3), Inches(4.85), Inches(9.4), Inches(0.55))
  cap.text_frame.word_wrap = True
  rr = cap.text_frame.paragraphs[0].add_run()
  rr.text = (
      'Left: masked + hubmatch thr=1 (MASK→commit).  '
      'Right: uniform ancestral (revise in gold window).  '
      'Teaser — decode stacks differ. GIF: docs/research/decode_viz/compare_…/preview.gif')
  rr.font.size = Pt(10)

  # 53 Conclusion
  s = prs.slides[53]
  h = find_text(s, 'Conclusion')
  if h:
    set_text(
        h,
        'Takeaways\n\n'
        'Conversion works (0.5 → 62 → 67 GSM).\n'
        'Components interact — shift helps; complementary alone hurts.\n'
        'Corruption is real; uniform needs ancestral decode.\n'
        'Keep the map explicit vs Fast-dLLM / BlockGen / Unifusion.')

  # 54 Thanks — keep
  # 55 Refs
  s = prs.slides[55]
  h = find_text(s, 'References')
  if h:
    set_text(h, 'References & ledger')
  for sh in shapes_text(s):
    if sh is h:
      continue
    if len(sh.text_frame.text) > 30:
      set_text(
          sh,
          'Fast-dLLM · BlockGen · Unifusion (design-space citations)\n\n'
          'Scoreboard: docs/research/PRESENTATION_SNAPSHOT_2026-09-18.md\n'
          'Decode GIF: docs/research/decode_viz/compare_…/preview.gif\n'
          'Claim K: docs/research/editable_tokens/')
      break

  # Curated order (original indices only — no clones)
  keep = [
      0,   # title
      1,   # toc
      3,   # research question
      2,   # design space
      6,   # three RQs
      7,   # related work section
      8,   # related work body
      10,  # paradigm vs components
      12,  # orthogonal layers
      20,  # ablations section
      21,  # ablation table
      15,  # component highlights
      14,  # paradigm results table
      36,  # corruption section
      37,  # corruption why
      46,  # decode viz
      16, 17, 18,  # caveats
      19,  # why factorization
      29,  # status board
      53,  # takeaways
      54,  # thanks
      55,  # refs
  ]
  keep_only(prs, keep)

  prs.save(str(OUT))
  shutil.copy2(OUT, ALSO)
  print('saved', OUT)
  print('slides', len(keep))

  prs2 = Presentation(str(OUT))
  print('verify open ok, n=', len(prs2.slides))
  for i, slide in enumerate(prs2.slides):
    pics = sum(1 for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    texts = []
    for sh in slide.shapes:
      if sh.has_text_frame:
        t = sh.text_frame.text.strip().split('\n')[0][:60]
        if t and 'Slide' not in t:
          texts.append(t)
    print(f'{i+1:02d} pics={pics} {texts[:2]}')


if __name__ == '__main__':
  build()
